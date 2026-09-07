"""
Trade Log Sheet & Data Export Module for VeteranDesk Streamlit Dashboard.

Generates structured, professional Excel (.xlsx) and CSV (.csv) exports joining:
- Trades (from Supabase 'trades' / 'demo_trades' table)
- Post-mortems and 4-verdict evaluations (from 'trade_journal' table)
- Independent discipline mistake flags (from 'mistake_audit_log' table)
- Double-entry ledger entries (from 'demo_ledger' table)
"""

from __future__ import annotations

from datetime import date, datetime
import io
from typing import Any, Dict, List, Optional, Tuple
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd

from veterandesk.database.session import db_manager
from veterandesk.logging import get_logger

logger = get_logger("veterandesk.dashboard.export")

TRADE_EXPORT_COLUMNS = [
    "Trade ID",
    "Ticker",
    "Date & Entry Time",
    "Signal Type",
    "Position Size (Shares)",
    "Entry Price (PKR)",
    "Stop Loss (PKR)",
    "Target Price (PKR)",
    "Exit Price (PKR)",
    "Exit Time",
    "Exit Reason",
    "Net PnL (PKR)",
    "Net PnL (%)",
    "Verdict",
    "Post-Mortem Analysis",
    "Transferable Lesson",
    "Mistake Flags",
    "Status",
]


def fetch_export_raw_data(
    client: Any = None,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[Dict[str, Any]]]:
    """
    Fetch read-only snapshot of trades, trade_journal, and mistake_audit_log from Supabase.
    """
    if client is None:
        client = db_manager.get_client()

    trades: List[Dict[str, Any]] = []
    try:
        res = client.table("trades").select("*").order("id", desc=True).limit(5000).execute()
        trades = res.data or []
    except Exception as ex:
        logger.warning("export_fetch_trades_failed", error=str(ex))

    # If trades table has no rows, check demo_trades fallback
    if not trades:
        try:
            res_demo = client.table("demo_trades").select("*").order("id", desc=True).limit(5000).execute()
            trades = res_demo.data or []
        except Exception as ex:
            logger.warning("export_fetch_demo_trades_failed", error=str(ex))

    journals: List[Dict[str, Any]] = []
    try:
        res_j = client.table("trade_journal").select("*").order("id", desc=True).limit(5000).execute()
        journals = res_j.data or []
    except Exception as ex:
        logger.warning("export_fetch_journals_failed", error=str(ex))

    mistakes: List[Dict[str, Any]] = []
    try:
        res_m = client.table("mistake_audit_log").select("*").order("id", desc=True).limit(5000).execute()
        mistakes = res_m.data or []
    except Exception as ex:
        logger.warning("export_fetch_mistakes_failed", error=str(ex))

    return trades, journals, mistakes


def build_trade_log_dataframe(
    trades: List[Dict[str, Any]],
    journals: List[Dict[str, Any]],
    mistakes: List[Dict[str, Any]],
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    ticker_filter: Optional[str] = None,
    verdict_filter: Optional[str] = None,
) -> pd.DataFrame:
    """
    Joins trades with journal verdicts/post-mortems and mistake audit flags into a flat DataFrame.
    Applies optional date range, ticker, and verdict filters.
    """
    if not trades:
        return pd.DataFrame(columns=TRADE_EXPORT_COLUMNS)

    # Index journals by trade_id
    journal_map: Dict[str, Dict[str, Any]] = {}
    for j in journals:
        tid = j.get("trade_id")
        if tid and tid not in journal_map:
            journal_map[tid] = j

    # Group mistakes by trade_id
    mistakes_map: Dict[str, List[str]] = {}
    for m in mistakes:
        tid = m.get("trade_id")
        if tid:
            severity = m.get("severity") or "CRITICAL"
            rule = m.get("rule_violated") or "DISCIPLINE_BREACH"
            details = m.get("details") or ""
            desc = f"[{severity}] {rule}"
            if details:
                desc += f": {details}"
            mistakes_map.setdefault(tid, []).append(desc)

    rows: List[Dict[str, Any]] = []

    for t in trades:
        tid = t.get("trade_id", "")
        ticker = str(t.get("ticker", "")).strip().upper()
        opened_at_raw = str(t.get("opened_at") or "")

        # Date Filtering (matching on entry date YYYY-MM-DD)
        trade_date_str = opened_at_raw[:10]
        if trade_date_str:
            try:
                trade_date = datetime.strptime(trade_date_str, "%Y-%m-%d").date()
                if start_date and trade_date < start_date:
                    continue
                if end_date and trade_date > end_date:
                    continue
            except ValueError:
                pass

        # Ticker Filtering
        if ticker_filter and ticker_filter != "All" and ticker != ticker_filter.upper():
            continue

        # Journal & Verdict lookup
        j = journal_map.get(tid, {})
        status = str(t.get("status") or "OPEN").upper()
        raw_verdict = j.get("verdict")
        if raw_verdict:
            verdict = str(raw_verdict)
        elif status == "CLOSED":
            verdict = "Pending"
        else:
            verdict = "Open"

        # Verdict Filtering
        if verdict_filter and verdict_filter != "All" and verdict != verdict_filter:
            continue

        entry_price = float(t.get("entry_price") or 0.0)
        shares = int(t.get("shares") or 0)
        cost_basis = entry_price * shares

        net_pnl = float(t["net_pnl"]) if t.get("net_pnl") is not None else None
        if net_pnl is not None and cost_basis > 0:
            pnl_pct = round((net_pnl / cost_basis) * 100.0, 2)
        else:
            pnl_pct = None

        trade_mistakes = mistakes_map.get(tid, [])
        mistake_flags = " | ".join(trade_mistakes) if trade_mistakes else "None"

        opened_at_clean = opened_at_raw[:19].replace("T", " ") if opened_at_raw else ""
        closed_at_raw = str(t.get("closed_at") or "")
        closed_at_clean = closed_at_raw[:19].replace("T", " ") if closed_at_raw else ""

        rows.append({
            "Trade ID": tid,
            "Ticker": ticker,
            "Date & Entry Time": opened_at_clean,
            "Signal Type": t.get("action", ""),
            "Position Size (Shares)": shares,
            "Entry Price (PKR)": round(entry_price, 2),
            "Stop Loss (PKR)": round(float(t.get("stop_loss") or 0.0), 2),
            "Target Price (PKR)": round(float(t.get("target_price") or 0.0), 2),
            "Exit Price (PKR)": round(float(t["exit_price"]), 2) if t.get("exit_price") is not None else None,
            "Exit Time": closed_at_clean,
            "Exit Reason": t.get("exit_reason") or "",
            "Net PnL (PKR)": round(net_pnl, 2) if net_pnl is not None else None,
            "Net PnL (%)": pnl_pct,
            "Verdict": verdict,
            "Post-Mortem Analysis": j.get("post_mortem_analysis") or "",
            "Transferable Lesson": j.get("transferable_lesson") or "",
            "Mistake Flags": mistake_flags,
            "Status": status,
        })

    if not rows:
        return pd.DataFrame(columns=TRADE_EXPORT_COLUMNS)

    return pd.DataFrame(rows)


def build_ledger_dataframe(ledger_entries: List[Dict[str, Any]]) -> pd.DataFrame:
    """Format raw demo_ledger records for Excel/CSV export."""
    if not ledger_entries:
        return pd.DataFrame(columns=["ID", "Transaction ID", "Trade ID", "Account Name", "Debit (PKR)", "Credit (PKR)", "Balance After (PKR)", "Description", "Timestamp"])
    
    rows = []
    for e in ledger_entries:
        created_at_clean = str(e.get("created_at") or "")[:19].replace("T", " ")
        rows.append({
            "ID": e.get("id"),
            "Transaction ID": e.get("transaction_id"),
            "Trade ID": e.get("trade_id") or "",
            "Account Name": e.get("account_name"),
            "Debit (PKR)": float(e.get("debit") or 0.0),
            "Credit (PKR)": float(e.get("credit") or 0.0),
            "Balance After (PKR)": float(e.get("balance_after") or 0.0),
            "Description": e.get("description"),
            "Timestamp": created_at_clean,
        })
    return pd.DataFrame(rows)


def build_mistakes_dataframe(mistakes: List[Dict[str, Any]]) -> pd.DataFrame:
    """Format raw mistake_audit_log records for Excel/CSV export."""
    if not mistakes:
        return pd.DataFrame(columns=["ID", "Trade ID", "Rule Violated", "Severity", "Details", "Detected At", "Acknowledged"])

    rows = []
    for m in mistakes:
        detected_at_clean = str(m.get("detected_at") or m.get("audited_at") or "")[:19].replace("T", " ")
        rows.append({
            "ID": m.get("id"),
            "Trade ID": m.get("trade_id") or "",
            "Rule Violated": m.get("rule_violated"),
            "Severity": m.get("severity"),
            "Details": m.get("details") or m.get("discrepancy_details"),
            "Detected At": detected_at_clean,
            "Acknowledged": bool(m.get("acknowledged", False)),
        })
    return pd.DataFrame(rows)


def generate_excel_export(
    df: pd.DataFrame,
    sheet_name: str = "Trade_Log",
) -> bytes:
    """
    Renders a pandas DataFrame into a styled, professional Excel workbook in-memory.
    Returns raw bytes suitable for st.download_button().
    """
    buffer = io.BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        # Freeze header row
        ws.freeze_panes = "A2"

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        currency_cols = {
            "Entry Price (PKR)",
            "Stop Loss (PKR)",
            "Target Price (PKR)",
            "Exit Price (PKR)",
            "Net PnL (PKR)",
            "Debit (PKR)",
            "Credit (PKR)",
            "Balance After (PKR)",
            "debit",
            "credit",
            "balance_after",
        }
        pct_cols = {"Net PnL (%)"}
        int_cols = {"Position Size (Shares)", "shares", "id", "ID"}
        wrap_cols = {
            "Post-Mortem Analysis",
            "Transferable Lesson",
            "Mistake Flags",
            "Discrepancy Details",
            "Details",
            "details",
            "Description",
            "description",
        }

        # Format header row
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        ws.row_dimensions[1].height = 28

        col_names = [cell.value for cell in ws[1]]

        # Format data rows
        for row in ws.iter_rows(min_row=2):
            for col_idx, cell in enumerate(row):
                cell.border = thin_border
                col_name = col_names[col_idx] if col_idx < len(col_names) else ""

                if col_name in currency_cols and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_name in pct_cols and isinstance(cell.value, (int, float)):
                    cell.number_format = '0.00"%"'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_name in int_cols and isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_name in wrap_cols:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    return buffer.getvalue()


def generate_csv_export(df: pd.DataFrame) -> bytes:
    """
    Renders a pandas DataFrame into UTF-8 CSV bytes in-memory.
    """
    return df.to_csv(index=False).encode("utf-8")
